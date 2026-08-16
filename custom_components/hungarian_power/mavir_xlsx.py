"""MAVIR XLSX response parser."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from .models import MavirChartData

LOCAL_TZ = ZoneInfo("Europe/Budapest")


class MavirParseError(ValueError):
    """Raised when MAVIR does not return a parseable chart workbook."""


def parse_mavir_xlsx(payload: bytes, chart_id: int) -> MavirChartData:
    """Parse the first worksheet in a MAVIR chart export."""
    if not payload.startswith(b"PK"):
        raise MavirParseError("MAVIR response is not an XLSX/OOXML workbook")

    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as err:
        raise MavirParseError(f"MAVIR chart {chart_id} workbook could not be opened") from err
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as err:
            raise MavirParseError("MAVIR workbook has no rows") from err

        headers = [str(value).strip() if value is not None else "" for value in header_row]
        timestamp_index = next(
            (
                index
                for index, header in enumerate(headers)
                if "időpont" in header.casefold() or "timestamp" in header.casefold()
            ),
            0,
        )

        latest_values: dict[str, float] = {}
        column_timestamps: dict[str, datetime] = {}
        latest_timestamp: datetime | None = None
        point_count = 0

        for row in rows:
            if not row:
                continue
            timestamp = _parse_timestamp(
                row[timestamp_index] if timestamp_index < len(row) else None
            )
            if timestamp is None:
                continue
            point_count += 1
            numeric_seen = False
            for index, header in enumerate(headers):
                if not header or index == timestamp_index or index >= len(row):
                    continue
                value = _parse_number(row[index])
                if value is None:
                    continue
                latest_values[header] = value
                column_timestamps[header] = timestamp
                numeric_seen = True
            if numeric_seen and (latest_timestamp is None or timestamp > latest_timestamp):
                latest_timestamp = timestamp

        if not latest_values:
            raise MavirParseError(f"MAVIR chart {chart_id} contains no numeric data")

        return MavirChartData(
            chart_id=chart_id,
            timestamp=latest_timestamp,
            columns=latest_values,
            column_timestamps=column_timestamps,
            point_count=point_count,
        )
    finally:
        workbook.close()


def _parse_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\xa0", " ").replace(" ", "")
    if not text or text in {"-", "—"}:
        return None
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return None


def _parse_timestamp(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=LOCAL_TZ)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time(), tzinfo=LOCAL_TZ)
    if value is None:
        return None

    text = " ".join(str(value).strip().split())
    for pattern in (
        "%Y.%m.%d %H:%M:%S %z",
        "%Y.%m.%d %H:%M %z",
        "%Y-%m-%d %H:%M:%S%z",
        "%Y-%m-%d %H:%M:%S %z",
    ):
        try:
            parsed = datetime.strptime(text, pattern)
        except ValueError:
            continue
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=LOCAL_TZ)
    return None
